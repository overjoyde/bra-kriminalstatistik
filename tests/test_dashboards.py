"""Bygger Excel- och HTML-dashboarden från fixturer (utdrag ur Brås och Domstolsverkets offentliga statistik).

Fixturerna i tests/fixtures/analysis är tolkade tabeller (utdata från parse_bra), så testet täcker
dashboardbyggena utan att behöva Brås Excel-filer eller nätet.
"""
import json
import re
import tempfile
import unittest
from pathlib import Path
from unittest import mock

try:
    import openpyxl
    import pandas as pd

    from brastat.analysis import excel_dashboard, html_dashboard
    from brastat.analysis import parse_bra as pb
except ImportError:  # pandas/openpyxl saknas
    pd = None

FIX = Path(__file__).resolve().parent / "fixtures" / "analysis"


def fixture_frames():
    m = pd.read_csv(FIX / "monthly.csv.gz")
    m["Datum"] = pd.to_datetime(dict(year=m["År"], month=m["Månad"], day=1))
    a = pd.read_csv(FIX / "annual.csv.gz")
    s = pd.read_csv(FIX / "suspects.csv.gz")
    s.columns = [str(c) for c in s.columns]
    b = pd.read_csv(FIX / "amnessidor_bedrageri.csv.gz", dtype=str)
    b["value"] = pd.to_numeric(b["value"])
    p = pd.read_csv(FIX / "amnessidor_pt.csv.gz", dtype=str)
    p["value"] = pd.to_numeric(p["value"])
    return m, a, s, (b, p)


@unittest.skipIf(pd is None, "pandas/openpyxl saknas")
class TestDashboards(unittest.TestCase):
    def setUp(self):
        m, a, s, am = fixture_frames()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        for name, value in (("monthly", m), ("annual", a), ("suspects", s), ("amnessidor", am)):
            patcher = mock.patch.object(pb, name, return_value=value)
            patcher.start()
            self.addCleanup(patcher.stop)

    def test_html_dashboard(self):
        out = Path(self.tmp.name) / "d.html"
        with mock.patch.object(html_dashboard, "domstat_data",
                               return_value=html_dashboard.domstat_data(str(FIX / "domstat_watchlist.csv"))):
            html_dashboard.build(str(out))
        html = out.read_text(encoding="utf-8")
        self.assertTrue(html.startswith("<!doctype html>"))
        self.assertNotIn("/*__DATA__*/", html)
        self.assertNotIn("/*__DARK__*/", html)
        m = re.search(r"const D = (.*?);\n", html)
        self.assertIsNotNone(m, "inbäddad data saknas")
        data = json.loads(m.group(1))
        for key in ("monthly", "annual", "suspects", "ntu", "ptcsv", "tm", "regions", "domstat"):
            self.assertIn(key, data)
        self.assertIn("Hela landet", data["monthly"])
        self.assertEqual(data["domstat"]["Inkomna brottmål"]["2025"], 143894.0)
        self.assertEqual(set(data["domstat"]), set(html_dashboard.DOMSTAT_PANEL))

    def test_html_without_domstat_hides_panel(self):
        out = Path(self.tmp.name) / "d.html"
        with mock.patch.object(html_dashboard, "domstat_data", return_value={}):
            html_dashboard.build(str(out))
        self.assertIn('"domstat":{}', out.read_text(encoding="utf-8"))

    def test_domstat_data_missing_file(self):
        self.assertEqual(html_dashboard.domstat_data(str(FIX / "finns-inte.csv")), {})

    def test_excel_dashboard(self):
        out = Path(self.tmp.name) / "d.xlsx"
        excel_dashboard.build(str(out))
        wb = openpyxl.load_workbook(out)
        self.assertIn("Dashboard", wb.sheetnames)
        self.assertTrue(any(n.startswith("Data_") for n in wb.sheetnames))
        data_sheets = [wb[n] for n in wb.sheetnames if n.startswith("Data_")]
        self.assertTrue(all(ws.max_row > 1 for ws in data_sheets))
        # KPI-cellerna är levande formler över datablad
        formulas = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("=")]
        self.assertGreater(len(formulas), 50)


if __name__ == "__main__":
    unittest.main()
