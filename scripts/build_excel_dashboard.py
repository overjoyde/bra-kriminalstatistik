"""Build the Excel dashboard dashboard/Bra_trendbevakning_dashboard.xlsx from raw Brå files.

Usage:
    python scripts/build_excel_dashboard.py
    (open and save once in Excel/LibreOffice so formula results are cached)

All KPI/trend cells are live formulas (SUMIFS) over the Data_* sheets, so the
workbook recalculates if the data sheets are replaced. Re-run this script after
`python scripts/fetch_tables.py` to pick up new months.
Source: Brottsförebyggande rådet (Brå), public statistics (PSI).
"""
from __future__ import annotations

import os
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

import parse_bra as pb

OUT = os.path.join(pb.ROOT, "dashboard", "Bra_trendbevakning_dashboard.xlsx")

FONT = "Arial"
NAVY = "1F3864"
F_TITLE = Font(name=FONT, size=16, bold=True, color=NAVY)
F_SUB = Font(name=FONT, size=10, italic=True, color="595959")
F_H = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_B = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_INPUT = Font(name=FONT, size=10, bold=True, color="0000FF")
F_SEC = Font(name=FONT, size=12, bold=True, color=NAVY)
FILL_H = PatternFill("solid", fgColor=NAVY)
FILL_BAND = PatternFill("solid", fgColor="F2F2F2")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_RED = PatternFill("solid", fgColor="F8CBAD")
FILL_GREEN = PatternFill("solid", fgColor="C6EFCE")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM = "#,##0"
PCT = "+0.0%;-0.0%;0.0%"
PCT_PLAIN = "0.0%"
DEC1 = "#,##0.0"

# Brottstyp/modus -> generell typologi (indirekt koppling, se docs/02-brottskoder-aml-ctf-bedrageri.md)
TYPOLOGY_MAP = {
    "Bedrägeri och annan oredlighet": "Kontext",
    "Social manipulation, totalt": "Social manipulation (APP)",
    "Befogenhetsbedrägeri": "Social manipulation (APP)",
    "Social manipulation av annan typ": "Social manipulation (APP)",
    "Romansbedrägeri": "Social manipulation (APP)",
    "Investeringsbedrägeri": "Social manipulation / företag",
    "Fakturabedrägeri": "Företag som brottsverktyg",
    "Identitetsbedrägeri": "Identitetsmissbruk",
    "Annonsbedrägeri": "Köp/sälj-bedrägeri",
    "Olovlig identitetsanvändning": "Identitetsmissbruk",
    "Kortbedrägeri utan fysiskt kort": "Kortbedrägeri (betalmetodsförflyttning)",
    "Penningtvättsbrott, totalt": "Penningtvätt / mule",
}
TM_MAP = TYPOLOGY_MAP  # bakåtkompatibelt namn
KPI_SERIES = list(TYPOLOGY_MAP)
REGION_ORDER = ["Stockholm", "Öst", "Väst", "Syd", "Mitt", "Bergslagen", "Nord"]

# Data sheet column layout (kept fixed so formulas can use whole-column ranges)
DM = "Data_manad"   # A År, B Månad, C Datum, D Region, E Serie, F Antal, G Källfil
DA = "Data_ar"      # A År, B Serie, C Antal, D Per 100 000 inv, E Källfil
DS = "Data_misstankta"
DN = "Data_amnessidor"  # A Källa, B Diagram, C Fråga, D Serie, E Etikett, F Värde


def style_header(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.border = F_H, FILL_H, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_body(ws, r1, r2, c1, c2, band=True):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BOX
            if cell.font is None or cell.font.name != FONT or cell.font == Font():
                cell.font = F_B
            if band and (r - r1) % 2 == 1:
                cell.fill = FILL_BAND


def change_format(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0.05"], fill=FILL_RED))
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["-0.05"], fill=FILL_GREEN))


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    ws["A2"] = sub
    ws["A2"].font = F_SUB
    ws.sheet_view.showGridLines = False


def write_df(ws, df, number_formats=None):
    ws.append(list(df.columns))
    style_header(ws, 1, 1, len(df.columns))
    for rec in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in rec])
    for c in range(1, len(df.columns) + 1):
        ws.column_dimensions[L(c)].width = 16
        for row in ws.iter_rows(min_row=2, min_col=c, max_col=c):
            row[0].font = F_B
            if number_formats and df.columns[c - 1] in number_formats:
                row[0].number_format = number_formats[df.columns[c - 1]]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def sumifs_m(serie_ref, region_ref, year_expr, month_crit):
    """Monthly data SUMIFS. month_crit is an Excel criterion expression, e.g. '"<="&$C$5'."""
    return (f"SUMIFS({DM}!$F:$F,{DM}!$E:$E,{serie_ref},{DM}!$D:$D,{region_ref},"
            f"{DM}!$A:$A,{year_expr},{DM}!$B:$B,{month_crit})")


def build():
    m = pb.monthly()
    a = pb.annual()
    s = pb.suspects()
    bed, pt = pb.amnessidor()

    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    names = ["Trender", "Årsutveckling", "Regioner", "Misstänkta", "Utsatthet NTU", "Om datat", DM, DA, DS, DN]
    sheets = {n: wb.create_sheet(n) for n in names}

    # ---------------- Data sheets ----------------
    dm = m[["År", "Månad", "Datum", "Region", "Serie", "Antal", "Källfil"]].copy()
    dm["Datum"] = dm["Datum"].dt.date
    write_df(sheets[DM], dm, {"Datum": "yyyy-mm", "Antal": NUM})
    write_df(sheets[DA], a[["År", "Serie", "Antal", "Per 100 000 inv", "Källfil"]], {"Antal": NUM, "Per 100 000 inv": DEC1})
    s_out = s.copy()
    write_df(sheets[DS], s_out, {c: NUM for c in pb.AGE_COLS})
    amn = pd.concat([
        bed.assign(Källa="Ämnessida Bedrägeri"),
        pt.assign(Källa="Ämnessida Penningtvätt/terrorfinansiering"),
    ])
    amn["label"] = amn["label"].map(lambda x: int(x) if str(x).isdigit() else x)
    amn = amn.rename(columns={"chart": "Diagram", "question": "Fråga", "series": "Serie", "label": "Etikett", "value": "Värde"})
    write_df(sheets[DN], amn[["Källa", "Diagram", "Fråga", "Serie", "Etikett", "Värde"]], {"Värde": "#,##0.0"})

    # ---------------- Dashboard ----------------
    title(ws, "Brå-trendbevakning – bedrägeri och penningtvätt",
          "Offentlig, aggregerad statistik från Brå. Anmälda brott, inte faktisk brottslighet. Brottskoderna anger inte betalmedel – kopplingen till en viss betaltjänst är indirekt via modus.")
    ws["A4"], ws["A5"], ws["A6"], ws["A7"] = "Senaste månad (preliminär)", "År", "Månad", "Senaste slutliga år"
    ws["C4"] = f"=MAX({DM}!$C:$C)"
    ws["C4"].number_format = "yyyy-mm"
    ws["C5"] = "=YEAR(C4)"
    ws["C6"] = "=MONTH(C4)"
    ws["C7"] = f"=MAX({DA}!$A:$A)"
    ws["C8"] = "Hela landet"
    ws["A8"] = "Region (KPI-tabellen)"
    for r in range(4, 9):
        ws.cell(row=r, column=1).font = F_BOLD
        ws.cell(row=r, column=3).font = F_B
        ws.cell(row=r, column=3).alignment = Alignment(horizontal="left")
    ws["C8"].font, ws["C8"].fill = F_INPUT, FILL_INPUT
    ws["D8"] = "← ändra till t.ex. Stockholm, Syd … (polisregioner finns från 2022). Datum/år räknas fram automatiskt ur datat."
    ws["D8"].font = F_SUB

    hdr = ["Brottstyp / modus", "Typologi",
           "Senaste månad", "Samma månad fg år", "Förändring",
           "Hittills i år", "Samma period fg år", "Förändring",
           "Andel mot äldre/ funktionsnedsatt (i år)",
           "Senaste slutliga år", "Per 100 000 inv", "Förändring 5 år"]
    HR = 10
    for i, h in enumerate(hdr, start=1):
        ws.cell(row=HR, column=i, value=h)
    style_header(ws, HR, 1, len(hdr))
    ws.row_dimensions[HR].height = 58
    for i, serie in enumerate(KPI_SERIES):
        r = HR + 1 + i
        ws.cell(row=r, column=1, value=serie)
        ws.cell(row=r, column=2, value=TM_MAP[serie])
        A = f"$A{r}"
        ws.cell(row=r, column=3, value="=" + sumifs_m(A, "$C$8", "$C$5", "$C$6"))
        ws.cell(row=r, column=4, value="=" + sumifs_m(A, "$C$8", "$C$5-1", "$C$6"))
        ws.cell(row=r, column=5, value=f'=IFERROR(C{r}/D{r}-1,"")')
        ws.cell(row=r, column=6, value="=" + sumifs_m(A, "$C$8", "$C$5", '"<="&$C$6'))
        ws.cell(row=r, column=7, value="=" + sumifs_m(A, "$C$8", "$C$5-1", '"<="&$C$6'))
        ws.cell(row=r, column=8, value=f'=IFERROR(F{r}/G{r}-1,"")')
        if serie in pb.ELDER_PARENTS:
            eld = f'{A}&"{pb.ELDER_SUFFIX}"'
            ws.cell(row=r, column=9, value=f'=IFERROR({sumifs_m(eld, "$C$8", "$C$5", chr(34) + "<=" + chr(34) + "&$C$6")}/F{r},"")')
        else:
            ws.cell(row=r, column=9, value="–")
        ws.cell(row=r, column=10, value=f"=SUMIFS({DA}!$C:$C,{DA}!$B:$B,{A},{DA}!$A:$A,$C$7)")
        ws.cell(row=r, column=11, value=f"=SUMIFS({DA}!$D:$D,{DA}!$B:$B,{A},{DA}!$A:$A,$C$7)")
        ws.cell(row=r, column=12, value=f'=IFERROR(J{r}/SUMIFS({DA}!$C:$C,{DA}!$B:$B,{A},{DA}!$A:$A,$C$7-5)-1,"")')
        for c, fmt in zip(range(3, 13), [NUM, NUM, PCT, NUM, NUM, PCT, PCT_PLAIN, NUM, DEC1, PCT]):
            ws.cell(row=r, column=c).number_format = fmt
    last = HR + len(KPI_SERIES)
    style_body(ws, HR + 1, last, 1, len(hdr))
    for r in range(HR + 1, last + 1):
        ws.cell(row=r, column=1).font = F_BOLD
        ws.cell(row=r, column=9).alignment = Alignment(horizontal="right")
    change_format(ws, f"E{HR+1}:E{last}")
    change_format(ws, f"H{HR+1}:H{last}")
    change_format(ws, f"L{HR+1}:L{last}")
    n = last + 1
    notes = [
        "Rött = ökning över 5 %, grönt = minskning över 5 %. Månadsjämförelser görs preliminärt mot preliminärt, så som Brå gör.",
        "OBS penningtvätt: den preliminära månadsstatistiken ligger långt under den slutliga (2025: ca 11 700 preliminärt mot 19 571 slutligt). Använd slutliga årssiffror för nivåer.",
        "Förändring 5 år har en seriebrytning för identitetsbedrägeri, som registreras annorlunda sedan 2022–2023. Se fliken 'Om datat'.",
        "Modus-uppdelningen (social manipulation m.m.) finns från 2019. 'Förändring 5 år' blir tom om serien saknas för jämförelseåret.",
    ]
    for i, t in enumerate(notes):
        ws.cell(row=n + i, column=1, value=t).font = F_SUB

    widths = [34, 26, 12, 12, 11, 12, 12, 11, 14, 12, 11, 11]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = f"A{HR+1}"

    # ---------------- Trender ----------------
    t = sheets["Trender"]
    title(t, "Månadstrender – hela landet (preliminär statistik)",
          "Antal anmälda brott per månad samt rullande 12 månader (R12). Formler mot Data_manad.")
    t_series = ["Bedrägeri och annan oredlighet", "Social manipulation, totalt", "Befogenhetsbedrägeri",
                "Social manipulation av annan typ", "Romansbedrägeri", "Investeringsbedrägeri", "Fakturabedrägeri",
                "Identitetsbedrägeri", "Annonsbedrägeri", "Kortbedrägeri utan fysiskt kort",
                "Olovlig identitetsanvändning", "Penningtvättsbrott, totalt"]
    TR = 4
    t.cell(row=TR, column=1, value="Månad")
    for j, sname in enumerate(t_series, start=2):
        t.cell(row=TR, column=j, value=sname)
    r12_start = 2 + len(t_series) + 1  # one blank column between blocks
    for j, sname in enumerate(t_series):
        t.cell(row=TR, column=r12_start + j, value=f"R12 {sname}")
    style_header(t, TR, 1, 1 + len(t_series))
    style_header(t, TR, r12_start, r12_start + len(t_series) - 1)
    t.row_dimensions[TR].height = 60
    nat = m[m.Region == "Hela landet"]
    months = sorted(d for d in nat.loc[nat["År"] >= pb.MODUS_FROM_YEAR, "Datum"].dt.date.unique())
    for i, d in enumerate(months):
        r = TR + 1 + i
        t.cell(row=r, column=1, value=d).number_format = "yyyy-mm"
        for j in range(len(t_series)):
            col = 2 + j
            t.cell(row=r, column=col,
                   value=f'=SUMIFS({DM}!$F:$F,{DM}!$E:$E,{L(col)}${TR},{DM}!$D:$D,"Hela landet",{DM}!$C:$C,$A{r})'
                   ).number_format = NUM
            rc = r12_start + j
            if i >= 11:
                t.cell(row=r, column=rc, value=f"=SUM({L(col)}{r-11}:{L(col)}{r})").number_format = NUM
    t_last = TR + len(months)
    style_body(t, TR + 1, t_last, 1, 1 + len(t_series), band=False)
    style_body(t, TR + 1, t_last, r12_start, r12_start + len(t_series) - 1, band=False)
    t.column_dimensions["A"].width = 10
    for c in range(2, r12_start + len(t_series)):
        t.column_dimensions[L(c)].width = 13
    t.freeze_panes = f"B{TR+1}"

    def line(ws_src, cols, title_, ytitle, first_row, last_row, width=24, height=9):
        ch = LineChart()
        ch.title, ch.y_axis.title, ch.height, ch.width = title_, ytitle, height, width
        ch.y_axis.number_format = "#,##0"
        ch.x_axis.number_format = "yyyy-mm"
        ch.x_axis.majorTimeUnit = "months"
        for c in cols:
            ch.series.append(Series(Reference(ws_src, min_col=c, min_row=first_row, max_row=last_row),
                                    title=ws_src.cell(row=TR, column=c).value))
        ch.set_categories(Reference(ws_src, min_col=1, min_row=first_row, max_row=last_row))
        for sr in ch.series:
            sr.smooth = False
            sr.graphicalProperties.line.width = 20000
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        return ch

    col_of = {sname: r12_start + j for j, sname in enumerate(t_series)}
    raw_col = {sname: 2 + j for j, sname in enumerate(t_series)}
    anchor_col = L(r12_start + len(t_series) + 1)
    c1 = line(t, [col_of["Social manipulation, totalt"], col_of["Kortbedrägeri utan fysiskt kort"], col_of["Annonsbedrägeri"]],
              "R12: social manipulation jämfört med kortbedrägeri och annonsbedrägeri", "Anmälda brott, R12", TR + 12, t_last)
    t.add_chart(c1, f"{anchor_col}4")
    c2 = line(t, [col_of[x] for x in ["Befogenhetsbedrägeri", "Social manipulation av annan typ", "Identitetsbedrägeri",
                                       "Fakturabedrägeri", "Investeringsbedrägeri", "Romansbedrägeri"]],
              "R12: social manipulation och företagsrelaterade modus", "Anmälda brott, R12", TR + 12, t_last)
    t.add_chart(c2, f"{anchor_col}23")
    c3 = line(t, [raw_col["Penningtvättsbrott, totalt"]],
              "Penningtvättsbrott per månad (preliminärt, underskattar nivån)", "Anmälda brott", TR + 1, t_last)
    t.add_chart(c3, f"{anchor_col}42")

    # Dashboard charts (reference Trender)
    d1 = line(t, [col_of["Social manipulation, totalt"], col_of["Befogenhetsbedrägeri"], col_of["Social manipulation av annan typ"]],
              "R12 social manipulation", "Anmälda brott, R12", TR + 12, t_last, width=17, height=8)
    ws.add_chart(d1, f"A{n + len(notes) + 1}")

    # ---------------- Årsutveckling ----------------
    y = sheets["Årsutveckling"]
    title(y, "Årsutveckling – slutlig statistik, hela landet",
          "Tabell 100 (anmälda brott) 2016 och framåt. Formler mot Data_ar. Modus-uppdelning från 2019.")
    years = sorted(a["År"].unique())
    YR = 4
    y.cell(row=YR, column=1, value="Brottstyp")
    for j, yr in enumerate(years, start=2):
        y.cell(row=YR, column=j, value=int(yr))
    style_header(y, YR, 1, 1 + len(years))
    ann_series = KPI_SERIES + [x + pb.ELDER_SUFFIX for x in ["Befogenhetsbedrägeri", "Social manipulation av annan typ", "Investeringsbedrägeri", "Romansbedrägeri"]]
    for i, sname in enumerate(ann_series):
        r = YR + 1 + i
        y.cell(row=r, column=1, value=sname).font = F_BOLD
        for j in range(len(years)):
            c = 2 + j
            y.cell(row=r, column=c, value=f'=IFERROR(1/(1/SUMIFS({DA}!$C:$C,{DA}!$B:$B,$A{r},{DA}!$A:$A,{L(c)}${YR})),"")').number_format = NUM
    y_last = YR + len(ann_series)
    style_body(y, YR + 1, y_last, 1, 1 + len(years))

    # per 100k block
    P = y_last + 3
    y.cell(row=P - 1, column=1, value="Per 100 000 invånare").font = F_SEC
    y.cell(row=P, column=1, value="Brottstyp")
    for j, yr in enumerate(years, start=2):
        y.cell(row=P, column=j, value=int(yr))
    style_header(y, P, 1, 1 + len(years))
    for i, sname in enumerate(KPI_SERIES):
        r = P + 1 + i
        y.cell(row=r, column=1, value=sname).font = F_BOLD
        for j in range(len(years)):
            c = 2 + j
            y.cell(row=r, column=c, value=f'=IFERROR(1/(1/SUMIFS({DA}!$D:$D,{DA}!$B:$B,$A{r},{DA}!$A:$A,{L(c)}${P})),"")').number_format = DEC1
    p_last = P + len(KPI_SERIES)
    style_body(y, P + 1, p_last, 1, 1 + len(years))

    # terror financing + PT split (ämnessidor)
    T = p_last + 3
    y.cell(row=T - 1, column=1, value="Penningtvätt och finansiering av terrorism (ämnessidor, anmälda brott)").font = F_SEC
    y.cell(row=T, column=1, value="Serie")
    pt_years = sorted(int(v) for v in pt["label"].unique() if str(v).isdigit())
    for j, yr in enumerate(pt_years, start=2):
        y.cell(row=T, column=j, value=yr)
    style_header(y, T, 1, 1 + len(pt_years))
    pt_series = list(dict.fromkeys(pt["series"]))
    for i, sname in enumerate(pt_series):
        r = T + 1 + i
        y.cell(row=r, column=1, value=sname).font = F_BOLD
        for j in range(len(pt_years)):
            c = 2 + j
            y.cell(row=r, column=c, value=(f'=IF(COUNTIFS({DN}!$D:$D,$A{r},{DN}!$E:$E,{L(c)}${T},{DN}!$F:$F,"<>")=0,"..",'
                                           f'SUMIFS({DN}!$F:$F,{DN}!$D:$D,$A{r},{DN}!$E:$E,{L(c)}${T}))')).number_format = NUM
            y.cell(row=r, column=c).alignment = Alignment(horizontal="right")
    t_last2 = T + len(pt_series)
    style_body(y, T + 1, t_last2, 1, 1 + len(pt_years))
    y.cell(row=t_last2 + 1, column=1, value="'..' = uppgift saknas i källan. 'Penningtvättsbrott' avser 3–4 §, alltså exklusive grovt brott, förseelse och näringspenningtvätt.").font = F_SUB
    y.column_dimensions["A"].width = 58
    for c in range(2, 2 + len(years)):
        y.column_dimensions[L(c)].width = 11

    row_of = {sname: YR + 1 + i for i, sname in enumerate(ann_series)}
    ycol0 = years.index(pb.MODUS_FROM_YEAR) + 2
    ch_modus = BarChart()
    ch_modus.type, ch_modus.grouping, ch_modus.overlap = "col", "stacked", 100
    ch_modus.title, ch_modus.height, ch_modus.width = "Social manipulation per modus (stackat)", 9, 22
    for x in ["Befogenhetsbedrägeri", "Social manipulation av annan typ", "Romansbedrägeri", "Investeringsbedrägeri"]:
        r = row_of[x]
        sr = Series(Reference(y, min_col=ycol0, max_col=1 + len(years), min_row=r), title=x)
        ch_modus.series.append(sr)
    ch_modus.set_categories(Reference(y, min_col=ycol0, max_col=1 + len(years), min_row=YR))
    ch_modus.y_axis.number_format = "#,##0"
    ch_modus.x_axis.delete = ch_modus.y_axis.delete = False
    y.add_chart(ch_modus, f"{L(len(years) + 3)}4")

    ch_pt = BarChart()
    ch_pt.type, ch_pt.title, ch_pt.height, ch_pt.width = "col", "Penningtvättsbrott, totalt – slutlig statistik", 9, 22
    ch_pt.series.append(Series(Reference(y, min_col=2, max_col=1 + len(years), min_row=row_of["Penningtvättsbrott, totalt"]), title="Penningtvättsbrott, totalt"))
    ch_pt.set_categories(Reference(y, min_col=2, max_col=1 + len(years), min_row=YR))
    ch_pt.y_axis.number_format = "#,##0"
    ch_pt.legend = None
    ch_pt.x_axis.delete = ch_pt.y_axis.delete = False
    y.add_chart(ch_pt, f"{L(len(years) + 3)}23")

    # Dashboard: PT annual chart
    d2 = BarChart()
    d2.type, d2.title, d2.height, d2.width = "col", "Penningtvättsbrott per år (slutlig statistik)", 8, 17
    d2.series.append(Series(Reference(y, min_col=2, max_col=1 + len(years), min_row=row_of["Penningtvättsbrott, totalt"]), title="Penningtvättsbrott"))
    d2.set_categories(Reference(y, min_col=2, max_col=1 + len(years), min_row=YR))
    d2.legend = None
    d2.y_axis.number_format = "#,##0"
    d2.x_axis.delete = d2.y_axis.delete = False
    ws.add_chart(d2, f"G{n + len(notes) + 1}")

    # ---------------- Regioner ----------------
    g = sheets["Regioner"]
    title(g, "Polisregioner – hittills i år jämfört med samma period föregående år (preliminärt)",
          "Styrs av År/Månad på fliken Dashboard. Polisregioner finns i datat från 2022.")
    reg_series = ["Bedrägeri och annan oredlighet", "Social manipulation, totalt", "Befogenhetsbedrägeri",
                  "Social manipulation av annan typ", "Investeringsbedrägeri", "Fakturabedrägeri", "Identitetsbedrägeri",
                  "Annonsbedrägeri", "Kortbedrägeri utan fysiskt kort", "Olovlig identitetsanvändning", "Penningtvättsbrott, totalt"]
    blocks = [("Antal hittills i år", "cur"), ("Förändring mot samma period föregående år", "chg"), ("Andel av riket (hittills i år)", "share")]
    rr = 4
    block_rows = {}
    for label, kind in blocks:
        g.cell(row=rr, column=1, value=label).font = F_SEC
        rr += 1
        g.cell(row=rr, column=1, value="Region")
        for j, sname in enumerate(reg_series, start=2):
            g.cell(row=rr, column=j, value=sname)
        style_header(g, rr, 1, 1 + len(reg_series))
        g.row_dimensions[rr].height = 58
        hrow = rr
        regions = REGION_ORDER + (["Hela landet"] if kind != "share" else [])
        for i, reg in enumerate(regions):
            r = hrow + 1 + i
            g.cell(row=r, column=1, value=reg).font = F_BOLD
            for j in range(len(reg_series)):
                c = 2 + j
                sref = f"{L(c)}${hrow}"
                cur = sumifs_m(sref, f"$A{r}", "Dashboard!$C$5", '"<="&Dashboard!$C$6')
                if kind == "cur":
                    v, fmt = "=" + cur, NUM
                elif kind == "chg":
                    prev = sumifs_m(sref, f"$A{r}", "Dashboard!$C$5-1", '"<="&Dashboard!$C$6')
                    v, fmt = f'=IFERROR({cur}/{prev}-1,"")', PCT
                else:
                    tot = sumifs_m(sref, '"Hela landet"', "Dashboard!$C$5", '"<="&Dashboard!$C$6')
                    v, fmt = f'=IFERROR({cur}/{tot},"")', PCT_PLAIN
                g.cell(row=r, column=c, value=v).number_format = fmt
        last_r = hrow + len(regions)
        style_body(g, hrow + 1, last_r, 1, 1 + len(reg_series))
        block_rows[kind] = (hrow, last_r)
        if kind == "chg":
            change_format(g, f"B{hrow+1}:{L(1+len(reg_series))}{last_r}")
        if kind == "share":
            g.conditional_formatting.add(f"B{hrow+1}:{L(1+len(reg_series))}{last_r}",
                                         ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="8EA9DB"))
        rr = last_r + 3
    g.cell(row=rr - 2, column=1, value="Summan av regionerna kan avvika något från hela landet (brott utan känd region).").font = F_SUB
    g.column_dimensions["A"].width = 16
    for c in range(2, 2 + len(reg_series)):
        g.column_dimensions[L(c)].width = 14
    h, lr = block_rows["chg"]
    chg = BarChart()
    chg.type, chg.title, chg.height, chg.width = "bar", "Förändring hittills i år – social manipulation och penningtvätt", 9, 20
    for sname in ["Social manipulation, totalt", "Penningtvättsbrott, totalt"]:
        c = 2 + reg_series.index(sname)
        chg.series.append(Series(Reference(g, min_col=c, min_row=h + 1, max_row=lr - 1), title=sname))
    chg.set_categories(Reference(g, min_col=1, min_row=h + 1, max_row=lr - 1))
    chg.x_axis.number_format = "0%"
    chg.y_axis.number_format = "0%"
    chg.x_axis.delete = chg.y_axis.delete = False
    g.add_chart(chg, f"{L(len(reg_series) + 3)}4")

    # ---------------- Misstänkta ----------------
    ms = sheets["Misstänkta"]
    title(ms, "Misstänkta personer efter ålder (tabell 220, slutlig statistik, hela landet)",
          "Kan antyda rekrytering av unga, t.ex. som målvakter eller för att upplåta konton (money mules). Jämför endast på aggregerad nivå.")
    s_years = sorted(s["År"].unique())
    groups = [("15–17 år", ["15", "16", "17"]), ("18–20 år", ["18", "19", "20"]), ("21–24 år", ["21-24"]),
              ("25–29 år", ["25-29"]), ("30–39 år", ["30-39"]), ("40 år och äldre", ["40-49", "50-59", "60-"])]
    col_letter = {k: L(3 + i) for i, k in enumerate(pb.AGE_COLS)}  # Data_misstankta: A År, B Serie, C.. age cols
    rr = 4
    share_rows = {}
    for sname in ["Penningtvättsbrott, totalt", "Social manipulation, totalt", "Befogenhetsbedrägeri", "Bedrägeri och annan oredlighet"]:
        ms.cell(row=rr, column=1, value=sname).font = F_SEC
        rr += 1
        hdrs = ["År", "Samtliga"] + [gname for gname, _ in groups] + ["Andel 15–24 år", "Andel kvinnor"]
        for j, hname in enumerate(hdrs, start=1):
            ms.cell(row=rr, column=j, value=hname)
        style_header(ms, rr, 1, len(hdrs))
        hrow = rr
        ms.cell(row=hrow, column=len(hdrs) + 1, value=sname).font = Font(name=FONT, size=1, color="FFFFFF")  # key for formulas
        key = f"${L(len(hdrs)+1)}${hrow}"
        yrs = [yv for yv in s_years if ((s["Serie"] == sname) & (s["År"] == yv)).any()]
        for i, yv in enumerate(yrs):
            r = hrow + 1 + i
            ms.cell(row=r, column=1, value=int(yv))

            def sf(col):
                return f"SUMIFS({DS}!${col}:${col},{DS}!$A:$A,$A{r},{DS}!$B:$B,{key})"
            ms.cell(row=r, column=2, value="=" + sf(col_letter["Samtliga"])).number_format = NUM
            for k, (gname, cols) in enumerate(groups):
                ms.cell(row=r, column=3 + k, value="=" + "+".join(sf(col_letter[c]) for c in cols)).number_format = NUM
            ms.cell(row=r, column=3 + len(groups), value=f'=IFERROR((C{r}+D{r}+E{r})/B{r},"")').number_format = PCT_PLAIN
            ms.cell(row=r, column=4 + len(groups), value=f'=IFERROR({sf(col_letter["Kvinnor"])}/B{r},"")').number_format = PCT_PLAIN
        last_r = hrow + len(yrs)
        style_body(ms, hrow + 1, last_r, 1, len(hdrs))
        share_rows[sname] = (hrow, last_r)
        rr = last_r + 3
    ms.cell(row=rr - 2, column=1, value="En person kan vara misstänkt för flera brottstyper. Summan per åldersgrupp exkluderar okänd ålder.").font = F_SUB
    ms.column_dimensions["A"].width = 30
    for c in range(2, 12):
        ms.column_dimensions[L(c)].width = 13
    yc = LineChart()
    yc.title, yc.height, yc.width = "Andel misstänkta 15–24 år", 9, 20
    yc.y_axis.number_format = "0%"
    share_col = 3 + len(groups)
    for sname in ["Penningtvättsbrott, totalt", "Social manipulation, totalt", "Bedrägeri och annan oredlighet"]:
        h, lr = share_rows[sname]
        sr = Series(Reference(ms, min_col=share_col, min_row=h + 1, max_row=lr), title=sname)
        yc.series.append(sr)
    h, lr = share_rows["Penningtvättsbrott, totalt"]
    yc.set_categories(Reference(ms, min_col=1, min_row=h + 1, max_row=lr))
    yc.x_axis.delete = yc.y_axis.delete = False
    ms.add_chart(yc, "M4")

    # ---------------- Utsatthet NTU ----------------
    u = sheets["Utsatthet NTU"]
    title(u, "Självrapporterad utsatthet (NTU) – andel av befolkningen 16–84 år",
          "Fångar även bedrägerier som inte anmälts. Källa: Nationella trygghetsundersökningen via Brås ämnessida Bedrägeri.")
    ntu_charts = ["Andel som utsatts för försäljningsbedrägeri", "Andel som utsatts för kort-/kreditbedrägeri"]
    ntu_years = sorted(int(v) for v in bed.loc[bed.chart.isin(ntu_charts), "label"].unique())
    UR = 4
    u.cell(row=UR, column=1, value="År")
    cols = []
    for ch_name in ntu_charts:
        for ser in ["Samtliga", "Män", "Kvinnor"]:
            cols.append((ch_name, ser))
    for j, (ch_name, ser) in enumerate(cols, start=2):
        u.cell(row=UR, column=j, value=f"{ch_name.replace('Andel som utsatts för ', '').capitalize()} – {ser}")
    style_header(u, UR, 1, 1 + len(cols))
    u.row_dimensions[UR].height = 45
    # hidden key rows for formulas
    for j, (ch_name, ser) in enumerate(cols, start=2):
        u.cell(row=2, column=10 + j, value=ch_name).font = Font(name=FONT, size=1, color="FFFFFF")
        u.cell(row=3, column=10 + j, value=ser).font = Font(name=FONT, size=1, color="FFFFFF")
    for i, yr in enumerate(ntu_years):
        r = UR + 1 + i
        u.cell(row=r, column=1, value=yr)
        for j in range(len(cols)):
            c = 2 + j
            kc = L(10 + c)
            u.cell(row=r, column=c, value=(f"=SUMIFS({DN}!$F:$F,{DN}!$B:$B,${kc}$2,{DN}!$D:$D,${kc}$3,{DN}!$E:$E,$A{r})/100")).number_format = PCT_PLAIN
    u_last = UR + len(ntu_years)
    style_body(u, UR + 1, u_last, 1, 1 + len(cols))
    u.column_dimensions["A"].width = 8
    for c in range(2, 2 + len(cols)):
        u.column_dimensions[L(c)].width = 16
    uc = LineChart()
    uc.title, uc.height, uc.width = "Andel utsatta (samtliga)", 9, 20
    uc.y_axis.number_format = "0%"
    uc.series.append(Series(Reference(u, min_col=2, min_row=UR + 1, max_row=u_last), title="Försäljningsbedrägeri"))
    uc.series.append(Series(Reference(u, min_col=5, min_row=UR + 1, max_row=u_last), title="Kort- och kreditbedrägeri"))
    uc.set_categories(Reference(u, min_col=1, min_row=UR + 1, max_row=u_last))
    uc.x_axis.delete = uc.y_axis.delete = False
    u.add_chart(uc, f"A{u_last + 3}")

    # ---------------- Om datat ----------------
    o = sheets["Om datat"]
    title(o, "Om datat, förbehåll och typologier", f"Arbetsboken genererades {date.today().isoformat()} av scripts/build_excel_dashboard.py")
    lines = [
        ("Källa", "Brottsförebyggande rådet (Brå). Fritt vidareutnyttjande (PSI), ange Brå som källa. Endast offentlig och aggregerad statistik, inga personuppgifter."),
        ("Månadsdata", "Tabell P4/P1x, anmälda brott per månad, preliminär statistik. Hela landet 2015–, polisregioner 2022–. Fil: raw/anmalda_brott/tidsserie_manad/."),
        ("Årsdata", "Tabell 100, slutlig statistik för hela landet 2016–, med antal och antal per 100 000 invånare. Fil: raw/anmalda_brott/ar_100_landet/."),
        ("Misstänkta", "Tabell 220, personer misstänkta efter brottstyp, ålder och kön, 2016–. Fil: raw/misstankta/220_brottstyp_alder_kon/."),
        ("NTU och ämnessidor", "Diagramdata från Brås ämnessidor Bedrägeri och Penningtvätt/terrorfinansiering. Fil: raw/amnessidor/."),
        ("Förbehåll 1", "Anmälda brott mäter anmälningar, inte faktisk brottslighet. Bedrägeri är ett seriebrott, så enskilda stora ärenden kan ge stora toppar."),
        ("Förbehåll 2", "Brottskoderna anger inte betalmedel. Kopplingen till en viss betaltjänst är alltid indirekt, via modus."),
        ("Förbehåll 3", "Preliminär månadsstatistik revideras i den slutliga årsstatistiken, i genomsnitt ca 6 % fler brott. För penningtvätt är skillnaden mycket större, eftersom brotten ofta registreras sent i utredningen."),
        ("Seriebrott", "Modus-uppdelningen för social manipulation, annons-, identitets- och kortbedrägeri m.fl. finns i datat från 2019. Identitetsbedrägeri sjunker kraftigt 2021–2023, delvis på grund av ändrad registrering. Läs docs/metod/ före tolkning."),
        ("Andel mot äldre", "Avser brott mot äldre eller funktionsnedsatta enligt Brås brottskod."),
        ("Tolkning av färger", "Rött betyder ökning över 5 % och grönt minskning över 5 %. Det är en beskrivning av förändringen, inte en bedömning av risk."),
        ("Typologi – social manipulation", "Befogenhets-, romans- och investeringsbedrägeri; penningflöden via målvaktskonton."),
        ("Typologi – företag som brottsverktyg", "Investerings- och fakturabedrägeri, identitetsbedrägeri via bolag. Se Brå-rapporten Företag som brottsverktyg (2025)."),
        ("Typologi – identitetsmissbruk", "Identitets- och annonsbedrägeri, olovlig identitetsanvändning, andel unga misstänkta."),
        ("Förflyttning mellan betalmetoder", "Kortbedrägeri utan fysiskt kort jämfört med social manipulation."),
        ("Uppdatera", "1) python scripts/fetch_tables.py  2) python scripts/build_excel_dashboard.py  3) Öppna och spara i Excel/LibreOffice så att formlerna räknas om."),
        ("Formler", "Alla tal på flikarna Dashboard, Trender, Årsutveckling, Regioner, Misstänkta och Utsatthet NTU är formler mot Data_*-flikarna. Den gula cellen på Dashboard (region) går att ändra."),
    ]
    for i, (k, v) in enumerate(lines, start=4):
        o.cell(row=i, column=1, value=k).font = F_BOLD
        cell = o.cell(row=i, column=2, value=v)
        cell.font = F_B
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        o.row_dimensions[i].height = 30
    o.column_dimensions["A"].width = 30
    o.column_dimensions["B"].width = 120

    for sh in wb.worksheets:
        sh.sheet_properties.tabColor = NAVY if sh.title in ("Dashboard",) else ("A6A6A6" if sh.title.startswith("Data_") else "8EA9DB")

    for sh in wb.worksheets:
        if not sh.title.startswith("Data_"):
            sh.page_setup.orientation = "landscape"
            sh.page_setup.fitToWidth, sh.page_setup.fitToHeight = 1, 0
            sh.sheet_properties.pageSetUpPr.fitToPage = True

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
